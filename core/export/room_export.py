"""部屋一覧の CSV / Excel 出力。

検出された部屋候補（確認済み・要確認・却下を含む）を、図面名・単位・
出力日時とともに書き出す。CSV は UTF-8 (BOM 付き) で Excel でも
そのまま開ける。Excel は openpyxl を使う。
"""
import csv
from datetime import datetime
from pathlib import Path
from typing import Iterable, List, Optional, Sequence

from core.models.room import (
    REVIEW_STATUS_LABELS,
    WALL_TYPE_LABELS,
    Room,
)

# 出力列（完成基準 §7: 部屋番号・部屋名・面積・状態・面積基準）。
COLUMNS = ["部屋番号", "部屋名", "面積 (m²)", "状態", "面積基準", "候補種別", "確信度"]


def _meta_rows(
    drawing_name: str, units: str, exported_at: datetime
) -> List[List[str]]:
    """出力ファイル先頭のメタ情報（項目/値の2列）。"""
    return [
        ["項目", "値"],
        ["図面", drawing_name],
        ["単位", units],
        ["出力日時", exported_at.strftime("%Y-%m-%d %H:%M:%S")],
    ]


def _room_values(room: Room) -> List[object]:
    """1 部屋分の出力値。面積・確信度は数値のまま返す。"""
    return [
        room.room_id,
        room.room_name,
        round(float(room.area_m2), 2),
        REVIEW_STATUS_LABELS.get(room.review_status, "要確認"),
        WALL_TYPE_LABELS.get(room.wall_type, "不明"),
        room.properties.get("candidate_type", ""),
        round(float(room.confidence), 2),
    ]


def _sorted_rooms(rooms: Iterable[Room]) -> List[Room]:
    return sorted(rooms, key=lambda r: r.room_id)


def export_rooms_csv(
    rooms: Iterable[Room],
    path: str | Path,
    *,
    drawing_name: str = "",
    units: str = "mm",
    exported_at: Optional[datetime] = None,
) -> Path:
    """部屋一覧を CSV (UTF-8 BOM 付き) で書き出す。"""
    path = Path(path)
    exported_at = exported_at or datetime.now()
    rows = _sorted_rooms(rooms)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(_meta_rows(drawing_name, units, exported_at))
        writer.writerow([])
        writer.writerow(COLUMNS)
        for room in rows:
            values = _room_values(room)
            values[2] = f"{values[2]:.2f}"
            values[6] = f"{values[6]:.2f}"
            writer.writerow(values)
    return path


def export_rooms_xlsx(
    rooms: Iterable[Room],
    path: str | Path,
    *,
    drawing_name: str = "",
    units: str = "mm",
    exported_at: Optional[datetime] = None,
) -> Path:
    """部屋一覧を Excel (.xlsx) で書き出す。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    path = Path(path)
    exported_at = exported_at or datetime.now()
    rows = _sorted_rooms(rooms)

    wb = Workbook()
    ws = wb.active
    ws.title = "部屋一覧"
    for row in _meta_rows(drawing_name, units, exported_at):
        ws.append(row)
    ws.append([])
    ws.append(COLUMNS)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for room in rows:
        ws.append(_room_values(room))
    # 面積と確信度は数値セルとして小数 2 桁表示にする。
    for row in ws.iter_rows(min_row=ws.max_row - len(rows) + 1, max_row=ws.max_row):
        row[2].number_format = "0.00"
        row[6].number_format = "0.00"
    widths = [18, 20, 14, 10, 10, 26, 10]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    wb.save(path)
    return path


def export_rooms(
    rooms: Iterable[Room],
    path: str | Path,
    *,
    drawing_name: str = "",
    units: str = "mm",
    exported_at: Optional[datetime] = None,
) -> Path:
    """拡張子 (.csv / .xlsx) で出力形式を選んで書き出す。"""
    path = Path(path)
    suffix = path.suffix.lower()
    kwargs = {
        "drawing_name": drawing_name,
        "units": units,
        "exported_at": exported_at,
    }
    if suffix == ".csv":
        return export_rooms_csv(rooms, path, **kwargs)
    if suffix == ".xlsx":
        return export_rooms_xlsx(rooms, path, **kwargs)
    raise ValueError(f"対応していない出力形式です: {path.suffix!r}（.csv / .xlsx のみ）")


__all__ = [
    "COLUMNS",
    "export_rooms",
    "export_rooms_csv",
    "export_rooms_xlsx",
]
