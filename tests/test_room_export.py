"""部屋一覧の CSV / Excel 出力テスト。"""
import csv
import os
from datetime import datetime

import pytest

from core.analyzer.room_detector import (
    associate_room_names,
    deduplicate_room_candidates,
    detect_closed_polyline_room_candidates,
    detect_linework_room_candidates,
)
from core.export.room_export import (
    COLUMNS,
    export_rooms,
    export_rooms_csv,
    export_rooms_xlsx,
)
from core.models.room import (
    REVIEW_STATUS_CONFIRMED,
    REVIEW_STATUS_REJECTED,
    Room,
)
from core.parser.dxf_parser import DxfParser

SAMPLES_DIR = os.path.join(os.path.dirname(__file__), "..", "samples")
FIXED_TIME = datetime(2026, 10, 26, 22, 30, 0)


def _sample_rooms(filename):
    """サンプル DXF から検出した部屋候補（名前関連付け済み）を返す。"""
    drawing = DxfParser().parse(os.path.join(SAMPLES_DIR, filename))
    polyline_rooms = detect_closed_polyline_room_candidates(drawing)
    linework_rooms = detect_linework_room_candidates(drawing)
    rooms = deduplicate_room_candidates(polyline_rooms, linework_rooms)
    return drawing, associate_room_names(drawing, rooms)


def test_csv_output_contains_meta_header_and_rooms(tmp_path):
    drawing, rooms = _sample_rooms("sample_01.dxf")
    assert len(rooms) == 1
    assert rooms[0].area_m2 == pytest.approx(15.0, rel=0.001)

    path = tmp_path / "rooms.csv"
    export_rooms_csv(
        rooms, path, drawing_name="sample_01.dxf",
        units=drawing.units, exported_at=FIXED_TIME,
    )

    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    assert rows[0] == ["項目", "値"]
    assert rows[1] == ["図面", "sample_01.dxf"]
    assert rows[2] == ["単位", "mm"]
    assert rows[3] == ["出力日時", "2026-10-26 22:30:00"]
    assert rows[4] == []
    assert rows[5] == COLUMNS
    assert rows[6][0] == "ROOM_CANDIDATE_000001"
    assert rows[6][2] == "15.00"
    # 検出直後はすべて「要確認」。
    assert rows[6][3] == "要確認"
    assert rows[6][4] == "不明"
    assert rows[6][5] == "closed_straight_polyline"


def test_xlsx_output_matches_csv_values(tmp_path):
    from openpyxl import load_workbook

    drawing, rooms = _sample_rooms("sample_02.dxf")
    assert len(rooms) == 2

    path = tmp_path / "rooms.xlsx"
    export_rooms_xlsx(
        rooms, path, drawing_name="sample_02.dxf",
        units=drawing.units, exported_at=FIXED_TIME,
    )

    wb = load_workbook(path)
    ws = wb["部屋一覧"]
    assert ws["A2"].value == "図面"
    assert ws["B2"].value == "sample_02.dxf"
    assert ws["B4"].value == "2026-10-26 22:30:00"
    assert [c.value for c in ws[6]] == COLUMNS
    data = [[c.value for c in row] for row in ws.iter_rows(min_row=7, max_row=8)]
    assert data[0][0] == "ROOM_CANDIDATE_000001"
    assert data[0][2] == pytest.approx(12.0)
    assert data[1][2] == pytest.approx(12.0)
    assert data[0][3] == "要確認"
    # 面積は数値セルとして保存される（openpyxl は整数値を int として読む）。


def test_review_status_is_exported(tmp_path):
    _, rooms = _sample_rooms("sample_02.dxf")
    rooms[0].review_status = REVIEW_STATUS_CONFIRMED
    rooms[1].review_status = REVIEW_STATUS_REJECTED

    path = tmp_path / "rooms.csv"
    export_rooms_csv(rooms, path, exported_at=FIXED_TIME)

    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    assert rows[6][3] == "確定"
    assert rows[7][3] == "却下"


def test_empty_room_list_exports_header_only(tmp_path):
    path = tmp_path / "empty.csv"
    export_rooms_csv([], path, drawing_name="empty.dxf", exported_at=FIXED_TIME)

    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    assert rows[5] == COLUMNS
    assert len(rows) == 6


def test_export_rooms_dispatches_by_suffix(tmp_path):
    _, rooms = _sample_rooms("sample_01.dxf")
    csv_path = export_rooms(rooms, tmp_path / "out.csv", exported_at=FIXED_TIME)
    xlsx_path = export_rooms(rooms, tmp_path / "out.xlsx", exported_at=FIXED_TIME)
    assert csv_path.suffix == ".csv"
    assert xlsx_path.suffix == ".xlsx"
    assert csv_path.exists()
    assert xlsx_path.exists()
    with pytest.raises(ValueError):
        export_rooms(rooms, tmp_path / "out.txt", exported_at=FIXED_TIME)


def test_room_review_status_defaults_and_validation():
    room = Room(room_id="R1", room_name="", area_m2=10.0)
    assert room.review_status == "to_confirm"
    room.review_status = REVIEW_STATUS_CONFIRMED
    assert room.review_status == "confirmed"
    assert room.properties["review_status"] == "confirmed"
    # to_dict / from_dict で確認状態が引き継がれる。
    restored = Room.from_dict(room.to_dict())
    assert restored.review_status == "confirmed"
    with pytest.raises(ValueError):
        room.review_status = "maybe"
    # 未知の保存値は「要確認」にフォールバックする。
    room.properties["review_status"] = "unknown_value"
    assert room.review_status == "to_confirm"
